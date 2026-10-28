"""Actual Forge extraction, Genesis sealing, and Spectra query integration."""
from __future__ import annotations
import hashlib,json,os,secrets,shutil,tempfile,unicodedata,uuid,zipfile
from contextlib import contextmanager
from pathlib import Path
from typing import Any
from .common import BoundaryError,atomic_bytes,atomic_json,canonical,hash_file,lock,now,private_dir,read_json,sha

PIPELINE='forge-source-blocks@1'

def initialize_keys(root:Path)->None:
    from axm_build.sign import hybrid1_keygen,hybrid1_public_key
    directory=private_dir(root/'keys')
    with lock(root,'genesis-keys'):
        private=directory/'publisher.secret';public=directory/'publisher.pub'
        if private.exists()!=public.exists():raise BoundaryError('Incomplete owner keypair; restore the original keys before proceeding.')
        if not private.exists():
            pub,secret=hybrid1_keygen();atomic_bytes(private,secret,replace=False);atomic_bytes(public,pub,replace=False)
        if hybrid1_public_key(private.read_bytes())!=public.read_bytes():raise BoundaryError('Owner signing keys do not match.')
        if not (directory/'spectra.secret').exists():atomic_bytes(directory/'spectra.secret',secrets.token_urlsafe(32).encode(),replace=False)

@contextmanager
def spectra_environment(root:Path):
    values={'SPECTRA_TRUSTED_PUBKEY':str(root/'keys/publisher.pub'),'SPECTRA_SYSTEM_KEY':(root/'keys/spectra.secret').read_text(),'SPECTRA_DEV_MODE':'0'}
    saved={k:os.environ.get(k) for k in values}
    try:
        os.environ.update(values);yield
    finally:
        for k,v in saved.items():
            if v is None:os.environ.pop(k,None)
            else:os.environ[k]=v

def preflight_source(path:Path,max_bytes:int=64*1024*1024)->None:
    if not path.is_file() or path.stat().st_size>max_bytes:raise BoundaryError('Source exceeds the 64 MiB per-file bound or is not a regular file.')
    if path.suffix.lower() in {'.docx','.xlsx','.pptx'}:
        with zipfile.ZipFile(path) as archive:
            entries=archive.infolist()
            if len(entries)>10000 or sum(x.file_size for x in entries)>128*1024*1024:raise BoundaryError('Office archive exceeds safe expansion limits.')
            if any('/externallinks/' in x.filename.lower() for x in entries):raise BoundaryError('External workbook links need an explicit adapter.')

def extract_document(path:Path)->dict:
    """Runs in a separately bounded process. Original bytes always stay intact."""
    from axm_forge.ingestion.extractors import extract
    from axm_build.common import normalize_source_text
    path=Path(path);preflight_source(path);suffix=path.suffix.lower();blocks=[];notes=[]
    if suffix in {'.csv','.xlsx'}:
        import csv
        if suffix=='.csv':
            with path.open(encoding='utf-8-sig',newline='') as f:
                for n,row in enumerate(csv.reader(f),1):
                    if n>50001:notes.append('Table extraction limited to 50,000 data rows.');break
                    if len(row)>100:raise BoundaryError('Table exceeds 100 columns.')
                    blocks.append({'text':' | '.join(row),'locator':{'kind':'row','row':n}})
        else:
            from openpyxl import load_workbook
            book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
            try:
                for sheet in book.worksheets:
                    for n,row in enumerate(sheet.iter_rows(values_only=True),1):
                        if n>50001:notes.append('Sheet '+sheet.title+' limited to 50,000 data rows.');break
                        if len(row)>100:raise BoundaryError('Sheet exceeds 100 columns.')
                        blocks.append({'text':' | '.join('' if x is None else str(x) for x in row),'locator':{'kind':'row','sheet':sheet.title,'row':n}})
            finally:book.close()
        extractor='openpyxl/csv source adapter; formulas preserved, not evaluated'
    elif suffix in {'.txt','.md','.text','.pdf','.docx','.html','.htm','.json'}:
        document=extract(path)
        blocks=[{'text':x.text,'locator':x.locator} for x in document.blocks]
        extractor='axm-core/Forge.extract'
    else:raise BoundaryError('Unsupported format; retained input needs a declared parser: '+suffix)
    out=[];total=0
    for block in blocks:
        text=normalize_source_text(str(block['text'])).strip('\n')
        if not text.strip():continue
        for offset in range(0,len(text),1800):
            piece=text[offset:offset+1800]
            total+=len(piece.encode())
            if total>2*1024*1024 or len(out)>=3000:
                notes.append('Extraction stopped at its text/block bound. Original remains complete.');break
            out.append({'text':piece,'locator':block['locator'],'extracted_character_offset':offset})
        if total>2*1024*1024 or len(out)>=3000:break
    if not out:raise BoundaryError('No extractable text. Scanned/image-only material needs a qualified OCR adapter.')
    return {'schema':'mary/extraction@1','pipeline':PIPELINE,'extractor':extractor,'format':suffix,'blocks':out,'notes':list(dict.fromkeys(notes)),'coverage':'partial' if notes else 'extracted','original_sha256':hash_file(path)}

def seal_document(root:Path,path:Path,extraction:dict,*,previous:str|None=None)->dict:
    from axm_build.compiler_generic import CompilerConfig,compile_generic_shard
    from axm_build.common import normalize_source_text
    from axm_verify.logic import verify_shard
    from blake3 import blake3
    initialize_keys(root)
    if hash_file(path)!=extraction['original_sha256']:raise BoundaryError('Source changed after extraction.')
    stage=private_dir(root/'shards'/('.stage-'+uuid.uuid4().hex))
    scratch=private_dir(root/'runtime'/('seal-'+uuid.uuid4().hex))
    try:
        spans=[];chunks=[];candidates=[];byte_pos=0
        for index,block in enumerate(extraction['blocks']):
            marker=f'[block {index:06d}]\n'
            section=normalize_source_text(marker+block['text'])
            marker_bytes=marker.encode();piece=section[len(marker):]
            start=byte_pos+len(marker_bytes);end=byte_pos+len(section.encode())
            spans.append({'id':f'b{index:06d}','byte_start':start,'byte_end':end,'text_sha256':sha(piece.encode()),'locator':block['locator'],'extracted_character_offset':block['extracted_character_offset']})
            candidates.append({'subject':f'document block {index:06d}','predicate':'contains','object':piece,'object_type':'literal:string','evidence':section,'tier':0})
            chunks.append(section);byte_pos+=len(section.encode())
        source=''.join(chunks);atomic_bytes(scratch/'source.txt',source.encode());atomic_bytes(scratch/'candidates.jsonl',b''.join(canonical(x)+b'\n' for x in candidates))
        original_name='original'+path.suffix.lower()
        binding={k:v for k,v in extraction.items() if k!='blocks'};binding.update(name=path.name,original_file=original_name,spans=spans,source_sha256=sha(source.encode()))
        atomic_json(scratch/'extraction.json',binding)
        cfg=CompilerConfig(source_path=scratch/'source.txt',candidates_path=scratch/'candidates.jsonl',out_dir=stage,private_key=(root/'keys/publisher.secret').read_bytes(),publisher_id='@mary-local-owner',publisher_name='Local workspace owner',namespace='mary/source',created_at=now(),title=path.name,license_spdx='LicenseRef-Private-Source',supersedes=(previous,) if previous else (),extra_content=((original_name,path),('extraction.json',scratch/'extraction.json')))
        if not compile_generic_shard(cfg):raise BoundaryError('Genesis rejected source compilation.')
        verification=verify_shard(stage,root/'keys/publisher.pub')
        if verification.get('status')!='PASS':raise BoundaryError('Independent Genesis verification failed: '+str(verification.get('errors')))
        shard_id='sh1_'+blake3((stage/'manifest.json').read_bytes()).hexdigest();target=root/'shards'/shard_id
        if target.exists():raise BoundaryError('Unexpected preexisting shard generation; inspect before replacing.')
        os.replace(stage,target)
        return {'shard_id':shard_id,'original_sha256':binding['original_sha256'],'name':path.name,'format':path.suffix.lower(),'coverage':binding['coverage'],'notes':binding['notes'],'span_count':len(spans),'bytes':sum(p.stat().st_size for p in target.rglob('*') if p.is_file()),'created_at':now()}
    finally:
        shutil.rmtree(scratch,ignore_errors=True)
        if stage.exists():shutil.rmtree(stage,ignore_errors=True)

def verified_document(root:Path,shard_id:str)->tuple[dict,bytes]:
    from axm_verify.logic import verify_shard
    from blake3 import blake3
    import re
    if not re.fullmatch(r'sh1_[a-f0-9]{64}',shard_id):raise BoundaryError('Invalid Genesis shard identity.')
    path=root/'shards'/shard_id
    if not path.is_dir() or path.is_symlink():raise BoundaryError('Source shard is unavailable.')
    if 'sh1_'+blake3((path/'manifest.json').read_bytes()).hexdigest()!=shard_id:raise BoundaryError('Shard identity drift.')
    result=verify_shard(path,root/'keys/publisher.pub')
    if result.get('status')!='PASS':raise BoundaryError('Source integrity failed; no evidence returned.')
    metadata=read_json(path/'content/extraction.json');data=(path/'content/source.txt').read_bytes()
    if sha(data)!=metadata['source_sha256']:raise BoundaryError('Extraction binding failed.')
    for span in metadata['spans']:
        text=data[span['byte_start']:span['byte_end']]
        if sha(text)!=span['text_sha256']:raise BoundaryError('Source span binding failed.')
    return metadata,data

def spectra_claims(root:Path,shards:list[str])->dict:
    from axiom_runtime.engine import SpectraEngine
    initialize_keys(root);temporary=private_dir(root/'runtime'/('spectra-'+uuid.uuid4().hex));engine=None
    try:
        with lock(root,'spectra-environment',timeout=20),spectra_environment(root):
            engine=SpectraEngine(audit_path=str(temporary/'audit.jsonl'),cache_path=str(temporary/'cache.jsonl'),db_path=str(temporary/'catalog.sqlite'),temp_root=str(temporary/'tmp'))
            mounted=[]
            for sid in shards:
                verified_document(root,sid)
                result=engine.mount(str(root/'shards'/sid),None,verify=True)
                if result.get('status')!='ok':raise BoundaryError('Spectra could not mount a verified source.')
                mounted.append(result['shard_id'])
            result=engine.query_json('SELECT claim_id, object, tier FROM claims')
            return {'mounted':mounted,'columns':result['columns'],'rows':result['rows']}
    finally:
        if engine is not None:
            try:engine.con.close()
            except Exception:pass
        shutil.rmtree(temporary,ignore_errors=True)
