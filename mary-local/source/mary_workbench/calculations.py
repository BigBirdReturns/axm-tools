"""Constrained SymPy and source-bound table operations. No generated code or SQL."""
from __future__ import annotations
import ast,csv,hashlib,io,json,math,re,zipfile
from decimal import Decimal,InvalidOperation,localcontext,Inexact,Overflow
from fractions import Fraction
from pathlib import Path
from typing import Any
import sympy as sp

NUM=re.compile(r'^[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?$')
GROUPED=re.compile(r'^[+-]?\d{1,3}(?:,\d{3})+(?:\.\d*)?(?:[eE][+-]?\d+)?$')
X=sp.Symbol('x',real=True)
FUNCTIONS={'sqrt':sp.sqrt,'sin':sp.sin,'cos':sp.cos,'tan':sp.tan,'log':sp.log,'exp':sp.exp,'abs':sp.Abs}

def decimal_value(value:Any)->Decimal:
    if isinstance(value,bool):raise ValueError('Boolean values are not numeric quantities.')
    text=str(value).strip()
    if GROUPED.fullmatch(text):text=text.replace(',','')
    if not NUM.fullmatch(text) or len(text)>120:raise ValueError('Selected nonempty value is not a bounded decimal number.')
    try:d=Decimal(text)
    except InvalidOperation as e:raise ValueError('Invalid decimal value.') from e
    if not d.is_finite() or len(d.as_tuple().digits)>80 or abs(d.adjusted())>300:raise ValueError('Numeric precision or exponent exceeds this task profile.')
    return d

def parse(source:str)->sp.Expr:
    try:tree=ast.parse(source,mode='eval')
    except (ValueError,SyntaxError) as e:raise ValueError('Invalid expression.') from e
    if sum(1 for _ in ast.walk(tree))>120:raise ValueError('Expression exceeds 120 syntax nodes.')
    def visit(node,depth=0):
        if depth>20:raise ValueError('Expression nesting exceeds 20.')
        if isinstance(node,ast.Constant):
            if type(node.value) not in (int,float):raise ValueError('Only numeric literals are accepted.')
            literal=ast.get_source_segment(source,node)
            if literal is None:raise ValueError('Numeric literal is unavailable.')
            d=decimal_value(literal);n,q=d.as_integer_ratio();return sp.Rational(n,q)
        if isinstance(node,ast.Name) and node.id=='x':return X
        if isinstance(node,ast.UnaryOp) and isinstance(node.op,(ast.USub,ast.UAdd)):
            value=visit(node.operand,depth+1);return -value if isinstance(node.op,ast.USub) else value
        if isinstance(node,ast.BinOp) and isinstance(node.op,(ast.Add,ast.Sub,ast.Mult,ast.Div,ast.Pow)):
            if isinstance(node.op,ast.Pow):
                right=node.right.operand if isinstance(node.right,ast.UnaryOp) and isinstance(node.right.op,(ast.UAdd,ast.USub)) else node.right
                if not isinstance(right,ast.Constant) or type(right.value) is not int or abs(right.value)>50:raise ValueError('Exponent must be a literal integer between -50 and 50.')
                if any(isinstance(n,ast.Pow) for n in ast.walk(node.left)):raise ValueError('Nested powers exceed this task profile.')
            left=visit(node.left,depth+1);right=visit(node.right,depth+1)
            if isinstance(node.op,ast.Add):return left+right
            if isinstance(node.op,ast.Sub):return left-right
            if isinstance(node.op,ast.Mult):return left*right
            if isinstance(node.op,ast.Div):
                if right==0:raise ValueError('Division by zero.')
                return left/right
            return left**right
        if isinstance(node,ast.Call) and isinstance(node.func,ast.Name) and node.func.id in FUNCTIONS:
            name=node.func.id
            if node.keywords or len(node.args) not in ({1,2} if name=='log' else {1}):raise ValueError('Unsupported function arguments.')
            if name=='exp' and any(isinstance(n,ast.Call) and isinstance(n.func,ast.Name) and n.func.id=='exp' for a in node.args for n in ast.walk(a)):raise ValueError('Nested exponentials exceed this task profile.')
            args=[visit(a,depth+1) for a in node.args]
            if name=='exp' and args[0].is_number and abs(args[0])>10000:raise ValueError('Exponential argument exceeds its bound.')
            return FUNCTIONS[name](*args)
        raise ValueError('Syntax is outside the declared mathematical operations.')
    result=visit(tree.body)
    if result.has(sp.oo,-sp.oo,sp.zoo,sp.nan):raise ValueError('Expression has an undefined or nonfinite result.')
    return result

def calculate(expression:str)->dict:
    if not isinstance(expression,str) or not 1<=len(expression.strip())<=1000:raise ValueError('Use an expression of 1–1,000 characters.')
    normalized=expression.strip().replace('^','**')
    if normalized.lower().startswith('solve:'):
        equation=normalized[6:].strip()
        if equation.count('=')!=1:raise ValueError('solve: requires one equation and one equals sign.')
        left,right=equation.split('=');result=sp.solveset(parse(left.strip())-parse(right.strip()),X,domain=sp.S.Reals)
        common={'kind':'equation','input':expression,'variable':'x','domain':'real'}
        if result is sp.S.EmptySet:return {**common,'solutions':[],'status':'resolved'}
        if not isinstance(result,sp.FiniteSet):return {**common,'solutions':[],'status':'conditional','condition':str(result)}
        return {**common,'solutions':[str(v) for v in sorted(result,key=sp.default_sort_key)],'status':'resolved'}
    result=parse(normalized)
    if result.is_Integer and len(str(result))>5000:raise ValueError('Exact integer result exceeds its display bound.')
    return {'kind':'calculation','input':expression,'exact':str(result),'decimal':str(result.evalf(15)),'free_symbols':sorted(str(x) for x in result.free_symbols),'domain':'complex' if result.is_real is False else 'real or symbolic'}

def digest(path:Path)->str:
    h=hashlib.sha256()
    with path.open('rb') as f:
        for b in iter(lambda:f.read(1024**2),b''):h.update(b)
    return h.hexdigest()

def read_table(path:Path,sheet:str|None)->tuple[list[str],list[list],str|None,list[dict]]:
    if not isinstance(path,Path) or not path.is_file() or path.is_symlink() or path.suffix.lower() not in {'.csv','.xlsx'} or path.stat().st_size>64*1024**2:raise ValueError('Use a regular CSV/XLSX source of at most 64 MiB.')
    rows=[];formulas=[];name=None
    if path.suffix.lower()=='.csv':
        with path.open(encoding='utf-8-sig',newline='') as f:
            for n,row in enumerate(csv.reader(f),1):
                if n>50001:raise ValueError('Table exceeds 50,000 data rows; no truncated calculation was returned.')
                if len(row)>100:raise ValueError('Table exceeds 100 columns.')
                rows.append(row)
    else:
        with zipfile.ZipFile(path) as z:
            entries=z.infolist()
            if len(entries)>10000 or sum(x.file_size for x in entries)>128*1024**2:raise ValueError('Workbook expansion exceeds its bound.')
            if any('vbaproject' in x.filename.lower() or '/externallinks/' in x.filename.lower() for x in entries):raise ValueError('Macros/external workbook links require a separate adapter.')
        from openpyxl import load_workbook
        book=load_workbook(path,read_only=True,data_only=False,keep_links=False)
        try:
            if sheet is not None and sheet not in book.sheetnames:raise ValueError('Requested sheet is absent.')
            ws=book[sheet] if sheet else book.worksheets[0];name=ws.title
            if ws.max_column and ws.max_column>100:raise ValueError('Sheet exceeds 100 columns.')
            for n,row in enumerate(ws.iter_rows(),1):
                if n>50001:raise ValueError('Sheet exceeds 50,000 data rows.')
                values=[]
                for i,cell in enumerate(row):
                    if i>=100:raise ValueError('Sheet exceeds 100 columns.')
                    values.append(cell.value)
                    if cell.data_type=='f':formulas.append({'row':n,'column_index':i,'cell':cell.coordinate,'formula':str(cell.value)})
                rows.append(values)
        finally:book.close()
    if not rows:raise ValueError('Table has no header.')
    headers=[str(x).strip() if x is not None else '' for x in rows[0]]
    if not headers or any(not x for x in headers) or len(set(headers))!=len(headers):raise ValueError('Headers must be nonempty and unique.')
    for n,row in enumerate(rows[1:],2):
        if len(row)!=len(headers):raise ValueError(f'Row {n} has a different number of cells from the header.')
    return headers,rows[1:],name,formulas

def analyze_table(path:Path,operation:str='describe',column:str|None=None,group_by:str|None=None,sheet:str|None=None)->dict:
    allowed={'describe','count','sum','mean','min','max','group_sum','group_count'}
    if operation not in allowed:raise ValueError('Unknown table operation.')
    headers,data,sheet,formulas=read_table(path,sheet)
    def missing(v):return v is None or isinstance(v,str) and not v.strip()
    result={'kind':'table','source_name':path.name,'source_sha256':digest(path),'format':path.suffix[1:].lower(),'sheet':sheet,'columns':headers,'row_count':len(data),'operation':operation,'column':column,'group_by':group_by,'metrics':{},'rows':[],'missing_rows':[],'formula_cells':formulas,'warnings':[]}
    if operation=='describe':
        result['metrics']={h:{'nonempty':sum(not missing(row[i]) for row in data),'missing':sum(missing(row[i]) for row in data),'formulas':sum(f['column_index']==i for f in formulas)} for i,h in enumerate(headers)};return result
    if column is None and operation=='count':result['metrics']={'count':len(data),'value':str(len(data))};return result
    if column not in headers:raise ValueError('Select an existing column.')
    ci=headers.index(column);numeric=operation in {'sum','mean','min','max','group_sum'}
    if numeric and any(f['column_index']==ci and f['row']>1 for f in formulas):raise ValueError('Selected column contains unevaluated formulas; aggregation is held rather than omitting those values.')
    selected=[]
    for rn,row in enumerate(data,2):
        value=row[ci]
        if missing(value):result['missing_rows'].append(rn);continue
        if numeric and isinstance(value,str) and value.startswith('='):raise ValueError('Formula-like selected cell cannot be treated as a number.')
        selected.append((rn,decimal_value(value) if numeric else value))
    if operation.startswith('group_'):
        if group_by not in headers:raise ValueError('Select an existing grouping column.')
        gi=headers.index(group_by);groups={}
        for rn,value in selected:
            group=data[rn-2][gi];label='(blank)' if missing(group) else str(group)
            groups.setdefault(label,[]).append(value)
            if len(groups)>1000:raise ValueError('Grouping exceeds 1,000 groups.')
        with localcontext() as ctx:
            ctx.prec=800;ctx.traps[Inexact]=True;ctx.traps[Overflow]=True
            result['rows']=[{'group':k,'value':str(sum(groups[k],Decimal(0))) if operation=='group_sum' else len(groups[k])} for k in sorted(groups)]
        return result
    if operation=='count':result['metrics']={'count':len(selected),'value':str(len(selected))};return result
    values=[v for _,v in selected]
    if not values:result['metrics']={'count':0,'value':None};result['warnings'].append('No numeric values were present.');return result
    with localcontext() as ctx:
        ctx.prec=800;ctx.traps[Inexact]=True;ctx.traps[Overflow]=True
        total=sum(values,Decimal(0))
        if operation=='sum':value=str(total)
        elif operation=='min':value=str(min(values))
        elif operation=='max':value=str(max(values))
        else:
            fraction=Fraction(total)/len(values);value=str(fraction)
    result['metrics']={'count':len(values),'value':value,'basis':'exact decimal' if operation!='mean' else 'exact rational'}
    if len(selected)<=1000:result['rows']=[{'row':r,'value':str(v)} for r,v in selected]
    else:result['warnings'].append('Detailed row provenance remains in the hash-identified original source; summary covers every selected row.')
    return result

def render_table_csv(result:dict)->str:
    if result.get('kind')!='table':raise ValueError('Expected a table result.')
    def safe(v):
        if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False)
        if isinstance(v,str) and not NUM.fullmatch(v) and (v.lstrip().startswith(('=','+','-','@')) or any(c in v for c in '\t\r\n')):return "'"+v
        return v
    stream=io.StringIO(newline='');writer=csv.writer(stream);rows=result.get('rows') or []
    if rows:
        columns=list(rows[0]);writer.writerow([safe(x) for x in columns])
        for row in rows:writer.writerow([safe(row.get(k,'')) for k in columns])
    else:
        writer.writerow(['metric','value'])
        for k,v in result.get('metrics',{}).items():writer.writerow([safe(k),safe(v)])
    return stream.getvalue()
